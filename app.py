from flask import Flask, render_template, request, redirect, url_for, flash, session, Response
from werkzeug.security import generate_password_hash, check_password_hash
from markupsafe import Markup
from datetime import datetime
from functools import wraps
import urllib.parse
import sqlite3
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import send_file

app = Flask(__name__)
app.secret_key = 'sgt_particular_secret_key_super_segura'
DB_NAME = 'sgt_particular.db'
PRECIO_PASAJE_DEFECTO = 5.00

# --- CONEXIÓN Y CREACIÓN DE BASE DE DATOS ---
def get_db_connection():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            rol TEXT NOT NULL,
            nombre TEXT NOT NULL,
            vehiculo_id INTEGER
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS vehiculos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            placa TEXT NOT NULL,
            modelo TEXT NOT NULL,
            tipo TEXT NOT NULL,
            capacidad INTEGER NOT NULL,
            chofer TEXT NOT NULL,
            estado TEXT NOT NULL,
            mantenimiento TEXT NOT NULL
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS rutas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            origen TEXT NOT NULL,
            destino TEXT NOT NULL,
            precio REAL NOT NULL
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS reservas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vehiculo_id INTEGER NOT NULL,
            nombre TEXT NOT NULL,
            telefono TEXT,
            asiento TEXT NOT NULL,
            hora_salida TEXT NOT NULL,
            fecha TEXT NOT NULL,
            ruta_id INTEGER NOT NULL,
            precio REAL NOT NULL,
            whatsapp_link TEXT
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS viaticos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vehiculo_id INTEGER NOT NULL,
            concepto TEXT NOT NULL,
            monto REAL NOT NULL,
            descripcion TEXT,
            fecha TEXT NOT NULL
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS historial_viajes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vehiculo TEXT NOT NULL,
            chofer TEXT NOT NULL,
            pasajeros_totales INTEGER NOT NULL,
            ingresos REAL NOT NULL,
            gastos REAL NOT NULL,
            ganancia_neta REAL NOT NULL,
            fecha_cierre TEXT NOT NULL
        )
    ''')

    cursor.execute('SELECT COUNT(*) FROM usuarios')
    if cursor.fetchone()[0] == 0:
        cursor.execute('''
            INSERT INTO usuarios (username, password, rol, nombre, vehiculo_id)
            VALUES (?, ?, ?, ?, ?)
        ''', ('admin', generate_password_hash('admin123'), 'admin', 'Administrador del Sistema', 1))

        cursor.execute('''
            INSERT INTO usuarios (username, password, rol, nombre, vehiculo_id)
            VALUES (?, ?, ?, ?, ?)
        ''', ('pedro', generate_password_hash('chofer123'), 'conductor', 'Pedro Picapiedra', 1))

    cursor.execute('SELECT COUNT(*) FROM vehiculos')
    if cursor.fetchone()[0] == 0:
        cursor.executemany('''
            INSERT INTO vehiculos (placa, modelo, tipo, capacidad, chofer, estado, mantenimiento)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', [
            ('PBC-1234', 'Hino Coaster', 'Buseta', 12, 'Pedro Picapiedra', 'En Servicio', 'Al día'),
            ('PBA-5678', 'Hyundai H350', 'Furgoneta', 15, 'Pablo Mármol', 'En Servicio', 'Al día'),
            ('PBT-9012', 'Mercedes-Benz Sprinter', 'Buseta', 19, 'Vilma Picapiedra', 'En Servicio', 'Al día')
        ])

    cursor.execute('SELECT COUNT(*) FROM rutas')
    if cursor.fetchone()[0] == 0:
        cursor.executemany('''
            INSERT INTO rutas (origen, destino, precio)
            VALUES (?, ?, ?)
        ''', [
            ('Puyo', 'Tena', 4.00),
            ('Puyo', 'Baños', 3.50),
            ('Puyo', 'Ambato', 6.00),
            ('Puyo', 'Macas', 5.00),
            ('Puyo', 'Quito', 25.00)
        ])

    conn.commit()
    conn.close()

init_db()

# --- DECORADORES DE SEGURIDAD ---
def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'usuario_id' not in session:
            flash('Por favor inicia sesión para acceder al sistema.', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'usuario_id' not in session or session.get('rol') != 'admin':
            flash('Acceso restringido únicamente a Administradores.', 'danger')
            return redirect(url_for('conductor'))
        return f(*args, **kwargs)
    return decorated_function

# --- AUTENTICACIÓN ---

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')

        conn = get_db_connection()
        user = conn.execute('SELECT * FROM usuarios WHERE LOWER(username) = LOWER(?)', (username,)).fetchone()
        conn.close()

        if user and check_password_hash(user['password'], password):
            session['usuario_id'] = user['id']
            session['username'] = user['username']
            session['nombre'] = user['nombre']
            session['rol'] = user['rol']
            session['vehiculo_id'] = int(user['vehiculo_id'] or 1)

            flash(f'¡Bienvenido {user["nombre"]}!', 'success')

            if user['rol'] == 'admin':
                return redirect(url_for('index'))
            else:
                return redirect(url_for('conductor'))
        else:
            flash('Usuario o contraseña incorrectos.', 'danger')

    return render_template('login.html')

@app.route('/registro', methods=['GET', 'POST'])
def registro():
    conn = get_db_connection()
    vehiculos = conn.execute('SELECT * FROM vehiculos').fetchall()

    if request.method == 'POST':
        nombre = request.form.get('nombre', '').strip()
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        vehiculo_id = int(request.form.get('vehiculo_id', 1))

        if not username or not password or not nombre:
            flash('Por favor completa todos los campos.', 'warning')
            conn.close()
            return redirect(url_for('registro'))

        usuario_existente = conn.execute('SELECT id FROM usuarios WHERE LOWER(username) = LOWER(?)', (username,)).fetchone()
        if usuario_existente:
            flash('El nombre de usuario ya existe.', 'danger')
            conn.close()
            return redirect(url_for('registro'))

        hashed_pw = generate_password_hash(password)
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO usuarios (username, password, rol, nombre, vehiculo_id)
            VALUES (?, ?, ?, ?, ?)
        ''', (username, hashed_pw, 'conductor', nombre, vehiculo_id))
        nuevo_id = cursor.lastrowid
        conn.commit()
        conn.close()

        session['usuario_id'] = nuevo_id
        session['username'] = username
        session['nombre'] = nombre
        session['rol'] = 'conductor'
        session['vehiculo_id'] = vehiculo_id

        flash('Registro e inicio de sesión exitoso.', 'success')
        return redirect(url_for('conductor'))

    conn.close()
    return render_template('registro.html', vehiculos=vehiculos)

@app.route('/logout')
def logout():
    session.clear()
    flash('Has cerrado sesión correctamente.', 'info')
    return redirect(url_for('login'))

# --- PANEL PRINCIPAL (ADMINISTRADOR) ---

@app.route('/')
@login_required
@admin_required
def index():
    vehiculo_id = request.args.get('vehiculo_id', default=1, type=int)
    search_query = request.args.get('buscar_pasajero', default='', type=str).strip().lower()

    conn = get_db_connection()
    flota = conn.execute('SELECT * FROM vehiculos').fetchall()
    unidad = conn.execute('SELECT * FROM vehiculos WHERE id = ?', (vehiculo_id,)).fetchone()
    
    if not unidad:
        unidad = flota[0]
        vehiculo_id = unidad['id']

    rutas = conn.execute('SELECT * FROM rutas').fetchall()

    if search_query:
        reservas_unidad = conn.execute('''
            SELECT * FROM reservas WHERE vehiculo_id = ? AND LOWER(nombre) LIKE ?
        ''', (vehiculo_id, f'%{search_query}%')).fetchall()
    else:
        reservas_unidad = conn.execute('SELECT * FROM reservas WHERE vehiculo_id = ?', (vehiculo_id,)).fetchall()

    asientos_ocupados = [r['asiento'] for r in reservas_unidad]
    stats = {
        'ocupados': len(asientos_ocupados),
        'disponibles': unidad['capacidad'] - len(asientos_ocupados)
    }

    viaticos_unidad = conn.execute('SELECT * FROM viaticos WHERE vehiculo_id = ?', (vehiculo_id,)).fetchall()
    total_viaticos = sum(v['monto'] for v in viaticos_unidad)
    total_ingresos = sum(r['precio'] for r in reservas_unidad)
    ganancia_neta = total_ingresos - total_viaticos

    conn.close()

    return render_template('index.html', 
                           flota=flota, 
                           unidad=unidad, 
                           reservas=reservas_unidad,
                           asientos_ocupados=asientos_ocupados,
                           stats=stats,
                           rutas=rutas,
                           viaticos=viaticos_unidad,
                           total_viaticos=total_viaticos,
                           total_ingresos=total_ingresos,
                           ganancia_neta=ganancia_neta,
                           search_query=search_query)

# --- PANEL CONDUCTOR ---

@app.route('/conductor')
@login_required
def conductor():
    if session.get('rol') == 'admin':
        vehiculo_id = request.args.get('vehiculo_id', type=int)
        if not vehiculo_id:
            vehiculo_id = session.get('vehiculo_id', 1)
    else:
        vehiculo_id = session.get('vehiculo_id', 1)

    conn = get_db_connection()
    unidad = conn.execute('SELECT * FROM vehiculos WHERE id = ?', (vehiculo_id,)).fetchone()
    if not unidad:
        unidad = conn.execute('SELECT * FROM vehiculos LIMIT 1').fetchone()

    reservas_unidad = conn.execute('SELECT * FROM reservas WHERE vehiculo_id = ?', (unidad['id'],)).fetchall()
    stats = {
        'ocupados': len(reservas_unidad),
        'disponibles': unidad['capacidad'] - len(reservas_unidad)
    }
    conn.close()

    return render_template('conductor.html', unidad=unidad, reservas=reservas_unidad, stats=stats)

# --- ACCIONES Y OPERACIONES ---

@app.route('/reservar', methods=['POST'])
@login_required
def reservar():
    vehiculo_id = int(request.form.get('vehiculo_id'))
    nombre = request.form.get('nombre')
    telefono = request.form.get('telefono', '').strip()
    asiento = request.form.get('asiento')
    hora_salida = request.form.get('hora_salida', 'Sin definir')
    ruta_id = int(request.form.get('ruta_id', 0))

    conn = get_db_connection()
    unidad = conn.execute('SELECT * FROM vehiculos WHERE id = ?', (vehiculo_id,)).fetchone()
    
    if unidad and unidad['estado'] == 'Mantenimiento':
        flash('El vehículo está en Mantenimiento. No se pueden realizar reservas.', 'danger')
        conn.close()
        return redirect(url_for('index', vehiculo_id=vehiculo_id))

    ruta_sel = conn.execute('SELECT * FROM rutas WHERE id = ?', (ruta_id,)).fetchone()
    precio_aplicado = ruta_sel['precio'] if ruta_sel else PRECIO_PASAJE_DEFECTO
    origen_destino = f"{ruta_sel['origen']} ➔ {ruta_sel['destino']}" if ruta_sel else "Ruta General"
    fecha_actual = datetime.now().strftime('%Y-%m-%d %H:%M')

    texto_whatsapp = (
        f"🚌 *BOLETO DIGITAL - SGT-Particular*\n\n"
        f"👤 *Pasajero:* {nombre}\n"
        f"💺 *Asiento:* {asiento}\n"
        f"📍 *Ruta:* {origen_destino}\n"
        f"⏰ *Hora de Salida:* {hora_salida}\n"
        f"💵 *Valor:* ${precio_aplicado:.2f}\n"
        f"🚐 *Unidad:* {unidad['modelo']} ({unidad['placa']})\n"
        f"📅 *Fecha:* {fecha_actual}\n\n"
        f"¡Gracias por viajar con nosotros!"
    )

    mensaje_url = urllib.parse.quote(texto_whatsapp)
    if telefono:
        whatsapp_link = f"https://api.whatsapp.com/send?phone={telefono}&text={mensaje_url}"
    else:
        whatsapp_link = f"https://api.whatsapp.com/send?text={mensaje_url}"

    conn.execute('''
        INSERT INTO reservas (vehiculo_id, nombre, telefono, asiento, hora_salida, fecha, ruta_id, precio, whatsapp_link)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (vehiculo_id, nombre, telefono, asiento, hora_salida, fecha_actual, ruta_id, precio_aplicado, whatsapp_link))
    conn.commit()
    conn.close()

    mensaje_html = Markup(
        f'Reserva registrada para {nombre} (Asiento {asiento}). '
        f'<a href="{whatsapp_link}" target="_blank" class="btn btn-sm btn-success ms-2 fw-bold">'
        f'<i class="bi bi-whatsapp me-1"></i>Enviar Boleto por WhatsApp</a>'
    )
    flash(mensaje_html, 'success')

    if session.get('rol') == 'conductor':
        return redirect(url_for('conductor'))
    return redirect(url_for('index', vehiculo_id=vehiculo_id))

@app.route('/eliminar_reserva/<int:reserva_id>', methods=['POST'])
@login_required
def eliminar_reserva(reserva_id):
    vehiculo_id = request.args.get('vehiculo_id', type=int, default=1)
    conn = get_db_connection()
    conn.execute('DELETE FROM reservas WHERE id = ?', (reserva_id,))
    conn.commit()
    conn.close()
    
    flash('Reserva cancelada correctamente.', 'warning')
    if session.get('rol') == 'conductor':
        return redirect(url_for('conductor'))
    return redirect(url_for('index', vehiculo_id=vehiculo_id))

@app.route('/agregar_vehiculo', methods=['POST'])
@login_required
@admin_required
def agregar_vehiculo():
    placa = request.form.get('placa')
    modelo = request.form.get('modelo')
    tipo = request.form.get('tipo', 'Buseta')
    capacidad = int(request.form.get('capacidad', 12))
    chofer = request.form.get('chofer')

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO vehiculos (placa, modelo, tipo, capacidad, chofer, estado, mantenimiento)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    ''', (placa, modelo, tipo, capacidad, chofer, 'En Servicio', 'Al día'))
    nuevo_id = cursor.lastrowid
    conn.commit()
    conn.close()

    flash(f'Vehículo {modelo} ({placa}) agregado.', 'success')
    return redirect(url_for('index', vehiculo_id=nuevo_id))

@app.route('/editar_vehiculo', methods=['POST'])
@login_required
@admin_required
def editar_vehiculo():
    vehiculo_id = int(request.form.get('vehiculo_id'))
    conn = get_db_connection()
    conn.execute('''
        UPDATE vehiculos 
        SET placa = ?, modelo = ?, tipo = ?, capacidad = ?, chofer = ?, mantenimiento = ?
        WHERE id = ?
    ''', (
        request.form.get('placa'),
        request.form.get('modelo'),
        request.form.get('tipo'),
        int(request.form.get('capacidad')),
        request.form.get('chofer'),
        request.form.get('mantenimiento'),
        vehiculo_id
    ))
    conn.commit()
    conn.close()
    flash('Vehículo actualizado correctamente.', 'success')
    return redirect(url_for('index', vehiculo_id=vehiculo_id))

@app.route('/agregar_ruta', methods=['POST'])
@login_required
@admin_required
def agregar_ruta():
    origen = request.form.get('origen', 'Puyo').strip()
    destino = request.form.get('destino').strip()
    precio = float(request.form.get('precio', 0.0))

    if destino and precio > 0:
        conn = get_db_connection()
        conn.execute('INSERT INTO rutas (origen, destino, precio) VALUES (?, ?, ?)', (origen, destino, precio))
        conn.commit()
        conn.close()
        flash(f'Ruta {origen} - {destino} (${precio:.2f}) guardada.', 'success')
    return redirect(url_for('index'))

@app.route('/editar_ruta', methods=['POST'])
@login_required
@admin_required
def editar_ruta():
    ruta_id = int(request.form.get('ruta_id'))
    conn = get_db_connection()
    conn.execute('''
        UPDATE rutas SET origen = ?, destino = ?, precio = ? WHERE id = ?
    ''', (request.form.get('origen'), request.form.get('destino'), float(request.form.get('precio')), ruta_id))
    conn.commit()
    conn.close()
    flash('Ruta actualizada.', 'info')
    return redirect(url_for('index'))

@app.route('/eliminar_ruta/<int:ruta_id>', methods=['POST'])
@login_required
@admin_required
def eliminar_ruta(ruta_id):
    conn = get_db_connection()
    conn.execute('DELETE FROM rutas WHERE id = ?', (ruta_id,))
    conn.commit()
    conn.close()
    flash('Ruta eliminada.', 'warning')
    return redirect(url_for('index'))

@app.route('/cambiar_estado_flota', methods=['POST'])
@login_required
def cambiar_estado_flota():
    vehiculo_id = int(request.form.get('vehiculo_id'))
    nuevo_estado = request.form.get('estado')
    
    conn = get_db_connection()
    conn.execute('UPDATE vehiculos SET estado = ? WHERE id = ?', (nuevo_estado, vehiculo_id))
    conn.commit()
    conn.close()

    flash(f'Estado actualizado a: {nuevo_estado}', 'info')
    if session.get('rol') == 'conductor':
        return redirect(url_for('conductor'))
    return redirect(url_for('index', vehiculo_id=vehiculo_id))

@app.route('/registrar_viatico', methods=['POST'])
@login_required
def registrar_viatico():
    vehiculo_id = int(request.form.get('vehiculo_id'))
    conn = get_db_connection()
    conn.execute('''
        INSERT INTO viaticos (vehiculo_id, concepto, monto, descripcion, fecha)
        VALUES (?, ?, ?, ?, ?)
    ''', (
        vehiculo_id,
        request.form.get('concepto'),
        float(request.form.get('monto', 0)),
        request.form.get('descripcion'),
        datetime.now().strftime('%Y-%m-%d %H:%M')
    ))
    conn.commit()
    conn.close()

    flash('Viático registrado.', 'success')
    if session.get('rol') == 'conductor':
        return redirect(url_for('conductor'))
    return redirect(url_for('index', vehiculo_id=vehiculo_id))

@app.route('/eliminar_viatico/<int:viatico_id>', methods=['POST'])
@login_required
@admin_required
def eliminar_viatico(viatico_id):
    conn = get_db_connection()
    v_item = conn.execute('SELECT vehiculo_id FROM viaticos WHERE id = ?', (viatico_id,)).fetchone()
    vehiculo_id = v_item['vehiculo_id'] if v_item else 1
    conn.execute('DELETE FROM viaticos WHERE id = ?', (viatico_id,))
    conn.commit()
    conn.close()

    flash('Registro de viático eliminado.', 'warning')
    return redirect(url_for('index', vehiculo_id=vehiculo_id))

@app.route('/liquidacion')
@login_required
@admin_required
def liquidacion():
    vehiculo_id = request.args.get('vehiculo_id', default=1, type=int)
    conn = get_db_connection()
    
    unidad = conn.execute('SELECT * FROM vehiculos WHERE id = ?', (vehiculo_id,)).fetchone()
    if not unidad:
        unidad = conn.execute('SELECT * FROM vehiculos LIMIT 1').fetchone()

    reservas_unidad = conn.execute('SELECT * FROM reservas WHERE vehiculo_id = ?', (unidad['id'],)).fetchall()
    viaticos_unidad = conn.execute('SELECT * FROM viaticos WHERE vehiculo_id = ?', (unidad['id'],)).fetchall()
    historial_viajes = conn.execute('SELECT * FROM historial_viajes ORDER BY id DESC').fetchall()

    total_ingresos = sum(r['precio'] for r in reservas_unidad)
    total_viaticos = sum(v['monto'] for v in viaticos_unidad)
    ganancia_neta = total_ingresos - total_viaticos
    conn.close()

    return render_template('liquidacion.html', 
                           unidad=unidad, 
                           reservas=reservas_unidad, 
                           viaticos=viaticos_unidad,
                           total_ingresos=total_ingresos,
                           total_viaticos=total_viaticos,
                           ganancia_neta=ganancia_neta,
                           historial_viajes=historial_viajes)

@app.route('/cerrar_caja', methods=['POST'])
@login_required
@admin_required
def cerrar_caja():
    vehiculo_id = int(request.form.get('vehiculo_id'))
    conn = get_db_connection()
    unidad = conn.execute('SELECT * FROM vehiculos WHERE id = ?', (vehiculo_id,)).fetchone()

    if unidad:
        reservas_unidad = conn.execute('SELECT * FROM reservas WHERE vehiculo_id = ?', (vehiculo_id,)).fetchall()
        viaticos_unidad = conn.execute('SELECT * FROM viaticos WHERE vehiculo_id = ?', (vehiculo_id,)).fetchall()

        total_ingresos = sum(r['precio'] for r in reservas_unidad)
        total_gastos = sum(v['monto'] for v in viaticos_unidad)
        ganancia_neta = total_ingresos - total_gastos

        conn.execute('''
            INSERT INTO historial_viajes (vehiculo, chofer, pasajeros_totales, ingresos, gastos, ganancia_neta, fecha_cierre)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (
            f"{unidad['modelo']} ({unidad['placa']})",
            unidad['chofer'],
            len(reservas_unidad),
            total_ingresos,
            total_gastos,
            ganancia_neta,
            datetime.now().strftime('%Y-%m-%d %H:%M')
        ))

        conn.execute('DELETE FROM reservas WHERE vehiculo_id = ?', (vehiculo_id,))
        conn.execute('DELETE FROM viaticos WHERE vehiculo_id = ?', (vehiculo_id,))
        conn.execute('UPDATE vehiculos SET estado = ? WHERE id = ?', ('En Servicio', vehiculo_id))
        conn.commit()

        flash(f'¡Cierre de caja completado! Ganancia neta: ${ganancia_neta:.2f}', 'success')

    conn.close()
    return redirect(url_for('liquidacion', vehiculo_id=vehiculo_id))

# --- EXPORTAR HISTORIAL DE CIERRES DE CAJA A EXCEL (.XLSX) ---
@app.route('/exportar_cierres_csv')
@login_required
@admin_required
def exportar_cierres_csv():
    conn = get_db_connection()
    cierres = conn.execute('SELECT * FROM historial_viajes ORDER BY id DESC').fetchall()
    conn.close()

    # Crear libro y hoja de cálculo de Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cierres de Caja"

    # Encabezados
    headers = ['ID', 'Vehículo', 'Chofer', 'Pasajeros Totales', 'Ingresos ($)', 'Gastos ($)', 'Ganancia Neta ($)', 'Fecha Cierre']
    ws.append(headers)

    # Estilos del encabezado
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Azul oscuro
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Aplicar estilos a la fila de encabezados
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Agregar filas de datos desde la BD
    for c in cierres:
        ws.append([
            c['id'],
            c['vehiculo'],
            c['chofer'],
            c['pasajeros_totales'],
            c['ingresos'],
            c['gastos'],
            c['ganancia_neta'],
            c['fecha_cierre']
        ])

    # Aplicar formato de moneda y alineación a las celdas de datos
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            
            # Formato de moneda $ para columnas 5, 6 y 7 (Ingresos, Gastos, Ganancia)
            if cell.column in [5, 6, 7]:
                cell.number_format = '"$"#,##0.00'
                cell.alignment = right_align
            elif cell.column in [1, 4, 8]:
                cell.alignment = center_align

    # Autoajustar el ancho de cada columna
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Guardar en memoria y retornar como descarga de Excel (.xlsx)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="reporte_cierres_caja.xlsx"
    )


# --- VISTA IMPRIMIBLE / PDF DE LISTA DE PASAJEROS ---
@app.route('/imprimir_pasajeros/<int:vehiculo_id>')
@login_required
def imprimir_pasajeros(vehiculo_id):
    conn = get_db_connection()
    unidad = conn.execute('SELECT * FROM vehiculos WHERE id = ?', (vehiculo_id,)).fetchone()
    reservas_unidad = conn.execute('SELECT * FROM reservas WHERE vehiculo_id = ? ORDER BY asiento ASC', (vehiculo_id,)).fetchall()
    conn.close()

    return render_template('imprimir_pasajeros.html', unidad=unidad, reservas=reservas_unidad, fecha_impresion=datetime.now().strftime('%Y-%m-%d %H:%M'))

if __name__ == '__main__':
    app.run(debug=True)
